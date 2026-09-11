"""S1 — ingest the raw workbook and run structural validation.

docs/architecture.md §4.2
"""
from __future__ import annotations

import hashlib

import openpyxl
import pandas as pd

import config
from schemas import validate_header, validate_raw_frame, validate_sheet_name


def file_sha256(path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def ingest() -> tuple[pd.DataFrame, list[str], str]:
    if not config.RAW_XLSX.exists():
        raise FileNotFoundError(
            f"Raw workbook not found at {config.RAW_XLSX}. "
            "It is intentionally excluded from version control (see .gitignore) "
            "and must be placed there manually before running the pipeline."
        )

    wb = openpyxl.load_workbook(config.RAW_XLSX, data_only=True)
    validate_sheet_name(wb.sheetnames)
    ws = wb[config.SHEET_NAME]

    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1, max_col=11))]
    validate_header(header)

    rows = [
        r
        for r in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=11, values_only=True)
        if not all(v is None or (isinstance(v, str) and v.strip() == "") for v in r)
    ]
    df = pd.DataFrame(rows, columns=config.RAW_COLUMNS)
    warnings = validate_raw_frame(df, config.EXPECTED_RAW_ROWS)
    return df, warnings, file_sha256(config.RAW_XLSX)
